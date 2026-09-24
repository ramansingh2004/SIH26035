"""Pure PDF/DOCX renderers over one immutable report plan."""

from __future__ import annotations

import html
import json
from dataclasses import dataclass
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass(frozen=True)
class ReportTable:
    title: str
    headers: tuple[str, ...]
    rows: tuple[tuple[str, ...], ...]


@dataclass(frozen=True)
class ReportSection:
    title: str
    paragraphs: tuple[str, ...] = ()
    tables: tuple[ReportTable, ...] = ()


@dataclass(frozen=True)
class ReportPlan:
    title: str
    status: str
    watermark: str | None
    sections: tuple[ReportSection, ...]


def _text(value: Any) -> str:
    if value is None:
        return "N/A"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(
            value,
            sort_keys=True,
            ensure_ascii=False,
            separators=(",", ":"),
        )
    return str(value)


def _mapping_rows(mapping: dict[str, Any], keys: tuple[str, ...]) -> tuple[tuple[str, str], ...]:
    return tuple((key.replace("_", " ").title(), _text(mapping.get(key))) for key in keys)


def build_plan(context: dict) -> ReportPlan:
    """Produce one semantic plan consumed identically by both renderers."""
    record = context["regulatory_record"]
    control = context["document_control"]
    session = record.get("session", {})
    lab = record.get("laboratory", {})
    manufacturer = record.get("manufacturer", {})
    instrument = record.get("instrument_master_at_approval", {})
    ruleset = record.get("ruleset_record", {})

    document_kind = context["document_kind"]
    preview = document_kind == "UNOFFICIAL_PREVIEW"
    report_number = control.get("report_number")

    cover = ReportSection(
        "Document Control",
        paragraphs=(
            "This document is an unofficial working preview and does not imply approval or issue."
            if preview
            else "This document was rendered from an immutable approved report context.",
        ),
        tables=(
            ReportTable(
                "Document control",
                ("Field", "Value"),
                (
                    ("Report Number", _text(report_number)),
                    ("Revision", _text(control.get("revision_no"))),
                    ("Report Status", _text(control.get("report_status"))),
                    ("Planned Issue Date", _text(control.get("planned_issue_date"))),
                    (
                        "Source Regulatory Revision",
                        _text(control.get("source_regulatory_revision")),
                    ),
                    ("Test Session ID", _text(session.get("id"))),
                    ("Application Number", _text(session.get("application_number"))),
                    ("Workflow Status", _text(session.get("workflow_status"))),
                    ("Evaluation Status", _text(session.get("evaluation_status"))),
                    ("Compliance Outcome", _text(session.get("compliance_outcome"))),
                ),
            ),
        ),
    )

    lab_section = ReportSection(
        "Laboratory",
        tables=(
            ReportTable(
                "Laboratory details",
                ("Field", "Value"),
                _mapping_rows(
                    lab,
                    (
                        "name",
                        "code",
                        "address_line1",
                        "address_line2",
                        "city",
                        "state",
                        "postal_code",
                        "country",
                        "phone",
                        "email",
                        "accreditation_no",
                    ),
                ),
            ),
        ),
    )

    manufacturer_section = ReportSection(
        "Manufacturer",
        tables=(
            ReportTable(
                "Manufacturer details",
                ("Field", "Value"),
                _mapping_rows(
                    manufacturer,
                    (
                        "name",
                        "registration_number",
                        "address_line1",
                        "address_line2",
                        "city",
                        "state",
                        "postal_code",
                        "country",
                        "contact_person",
                        "email",
                        "phone",
                    ),
                ),
            ),
        ),
    )

    instrument_section = ReportSection(
        "Instrument Identification",
        tables=(
            ReportTable(
                "Instrument master facts captured in the regulatory record",
                ("Field", "Value"),
                tuple(
                    (str(key).replace("_", " ").title(), _text(value))
                    for key, value in sorted(instrument.items())
                    if key
                    not in {
                        "created_at",
                        "updated_at",
                        "metadata_json",
                    }
                ),
            ),
            ReportTable(
                "Instrument ranges",
                ("Range", "Min", "Max", "d", "e"),
                tuple(
                    (
                        _text(row.get("range_no")),
                        _text(row.get("min_capacity_g")),
                        _text(row.get("max_capacity_g")),
                        _text(row.get("scale_interval_d_g")),
                        _text(row.get("verification_interval_e_g")),
                    )
                    for row in record.get("instrument_ranges_at_approval", [])
                ),
            ),
        ),
    )

    rules_section = ReportSection(
        "Ruleset Declaration",
        tables=(
            ReportTable(
                "Pinned ruleset",
                ("Field", "Value"),
                _mapping_rows(
                    ruleset,
                    (
                        "standard_code",
                        "edition",
                        "version",
                        "verification_status",
                        "configuration_hash",
                        "source_reference",
                    ),
                ),
            ),
        ),
    )

    summary_rows = []
    for row in sorted(record.get("sections", []), key=lambda item: item.get("section_number", 0)):
        summary_rows.append(
            (
                _text(row.get("section_number")),
                _text(row.get("section_name") or row.get("name") or row.get("section_code")),
                _text(row.get("applicability_status")),
                _text(row.get("evaluation_status")),
                _text(row.get("compliance_outcome")),
                _text(row.get("applicability_reason")),
            )
        )
    summary_section = ReportSection(
        "Executive Summary - All 17 Sections",
        tables=(
            ReportTable(
                "Stored section decisions",
                (
                    "No.",
                    "Section",
                    "Applicability",
                    "Evaluation",
                    "Outcome",
                    "Remarks",
                ),
                tuple(summary_rows),
            ),
        ),
    )

    equipment_rows = []
    for link in record.get("equipment_links", []):
        snap = link.get("equipment_snapshot") or {}
        equipment_rows.append(
            (
                _text(snap.get("category")),
                _text(snap.get("manufacturer")),
                _text(snap.get("model")),
                _text(snap.get("serial_number")),
                _text(snap.get("reference_number")),
                _text(snap.get("calibration_certificate_no")),
                _text(snap.get("calibration_date")),
                _text(snap.get("calibration_due_date")),
                _text(snap.get("accuracy_or_class")),
            )
        )
    equipment_section = ReportSection(
        "Test Equipment",
        tables=(
            ReportTable(
                "Equipment snapshots",
                (
                    "Type",
                    "Manufacturer",
                    "Model",
                    "Serial",
                    "Reference",
                    "Certificate",
                    "Cal. Date",
                    "Cal. Due",
                    "Accuracy/Class",
                ),
                tuple(equipment_rows),
            ),
        ),
    )

    environment_rows = tuple(
        (
            _text(row.get("measured_at")),
            _text(row.get("phase")),
            _text(row.get("temperature_c")),
            _text(row.get("relative_humidity_percent")),
            _text(row.get("pressure_hpa")),
            _text(row.get("voltage_v")),
            _text(row.get("remarks")),
        )
        for row in record.get("environment_readings", [])
    )
    environment_section = ReportSection(
        "Environmental Conditions",
        tables=(
            ReportTable(
                "Stored environmental trace",
                (
                    "Date/Time",
                    "Phase",
                    "Temperature",
                    "RH %",
                    "Pressure",
                    "Voltage",
                    "Remarks",
                ),
                environment_rows,
            ),
        ),
    )

    runs = record.get("runs", [])
    observations = record.get("observations", [])
    results = record.get("results", [])
    run_tables = []
    for run in runs:
        run_id = run.get("id")
        run_tables.append(
            ReportTable(
                f"Run {run.get('run_no')} -{run.get('test_code') or run.get('test_definition_id')}",
                ("Field", "Stored Value"),
                (
                    ("Run ID", _text(run_id)),
                    ("Procedure Context", _text(run.get("procedure_context"))),
                    ("Evaluation Status", _text(run.get("evaluation_status"))),
                    ("Compliance Outcome", _text(run.get("compliance_outcome"))),
                    ("Input Revision", _text(run.get("input_revision"))),
                    ("Current Result ID", _text(run.get("current_result_id"))),
                    (
                        "Observations",
                        _text([row for row in observations if row.get("test_run_id") == run_id]),
                    ),
                    (
                        "Results / Calculation Trace",
                        _text([row for row in results if row.get("test_run_id") == run_id]),
                    ),
                ),
            )
        )
    tests_section = ReportSection(
        "Test Runs, Observations and Stored Results",
        paragraphs=(
            "Values and decisions below are rendered exactly from stored snapshots; "
            "the report renderer performs no compliance calculation.",
        ),
        tables=tuple(run_tables),
    )

    construction = record.get("construction", {})
    construction_section = ReportSection(
        "Construction Examination",
        tables=(
            ReportTable(
                "Construction examination",
                ("Field", "Stored Value"),
                tuple(
                    (str(key).replace("_", " ").title(), _text(value))
                    for key, value in sorted((construction.get("examination") or {}).items())
                ),
            ),
            ReportTable(
                "Construction items",
                ("Item", "Stored Record"),
                tuple(
                    (
                        _text(row.get("requirement_key") or row.get("id")),
                        _text(row),
                    )
                    for row in construction.get("items", [])
                ),
            ),
        ),
    )

    checklist_section = ReportSection(
        "Checklist",
        tables=(
            ReportTable(
                "Versioned checklist wording and response",
                ("Requirement", "Wording", "Response", "Remarks"),
                tuple(
                    (
                        _text(item.get("rule", {}).get("requirement_key")),
                        _text(item.get("rule", {}).get("display_text")),
                        _text(item.get("response", {}).get("result")),
                        _text(item.get("response", {}).get("remarks")),
                    )
                    for item in record.get("checklist", [])
                ),
            ),
        ),
    )

    evidence_section = ReportSection(
        "Evidence Register",
        tables=(
            ReportTable(
                "Immutable evidence identities",
                (
                    "Target",
                    "Purpose",
                    "File",
                    "SHA-256",
                    "Object Version",
                ),
                tuple(
                    (
                        (
                            _text(item.get("link", {}).get("entity_type"))
                            + ":"
                            + _text(item.get("link", {}).get("entity_id"))
                        ),
                        _text(item.get("link", {}).get("purpose")),
                        _text(item.get("attachment", {}).get("file_name")),
                        _text(item.get("attachment", {}).get("sha256")),
                        _text(item.get("attachment", {}).get("object_version")),
                    )
                    for item in record.get("evidence", [])
                ),
            ),
        ),
    )

    approvals_section = ReportSection(
        "Review and Approval History",
        tables=(
            ReportTable(
                "Append-only approval actions",
                (
                    "Stage",
                    "Decision",
                    "Actor",
                    "Regulatory Revision",
                    "Timestamp",
                    "Reason/Comment",
                ),
                tuple(
                    (
                        _text(row.get("stage")),
                        _text(row.get("decision")),
                        _text(row.get("actor_id")),
                        _text(row.get("regulatory_revision")),
                        _text(row.get("created_at")),
                        _text(row.get("reason") or row.get("comment")),
                    )
                    for row in record.get("approval_actions", [])
                ),
            ),
        ),
    )

    integrity_section = ReportSection(
        "Reproducibility Manifest",
        tables=(
            ReportTable(
                "Renderer and source identities",
                ("Field", "Value"),
                (
                    ("Template Version", _text(context.get("template_version"))),
                    ("Renderer Manifest", _text(context.get("renderer_manifest"))),
                    (
                        "Ruleset Snapshot",
                        _text(record.get("ruleset_snapshot")),
                    ),
                ),
            ),
        ),
    )

    return ReportPlan(
        title=("OIML R 76 - Non-Automatic Weighing Instrument Type-Evaluation Test Report"),
        status="UNOFFICIAL PREVIEW" if preview else _text(control.get("report_status")),
        watermark="UNOFFICIAL PREVIEW" if preview else None,
        sections=(
            cover,
            lab_section,
            manufacturer_section,
            instrument_section,
            rules_section,
            summary_section,
            equipment_section,
            environment_section,
            tests_section,
            construction_section,
            checklist_section,
            evidence_section,
            approvals_section,
            integrity_section,
        ),
    )


def _cell_chunks(value: str, limit: int = 1200) -> tuple[str, ...]:
    if len(value) <= limit:
        return (value,)
    return tuple(value[index : index + limit] for index in range(0, len(value), limit))


def _expanded_rows(report_table: ReportTable) -> tuple[tuple[str, ...], ...]:
    expanded = []
    for row in report_table.rows:
        chunks = [_cell_chunks(cell) for cell in row]
        height = max(len(parts) for parts in chunks)
        for index in range(height):
            expanded.append(tuple(parts[index] if index < len(parts) else "" for parts in chunks))
    return tuple(expanded)


def _pdf_paragraph(value: str, style):
    return Paragraph(html.escape(value).replace("\n", "<br/>"), style)


def render_pdf(plan: ReportPlan) -> bytes:
    buffer = BytesIO()
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "ReportBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        spaceAfter=4,
    )
    table_text = ParagraphStyle(
        "ReportTableText",
        parent=body,
        fontSize=6,
        leading=7,
        wordWrap="CJK",
    )
    heading = ParagraphStyle(
        "ReportHeading",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=16,
        spaceBefore=8,
        spaceAfter=6,
    )
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        alignment=TA_CENTER,
    )

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=plan.title,
        author="SIH26035",
    )

    def page(canvas, document):
        canvas.saveState()
        width, height = A4
        canvas.setFont("Helvetica", 7)
        canvas.setFillGray(0.35)
        canvas.drawCentredString(width / 2, 8 * mm, f"Page {document.page}")
        if plan.watermark:
            canvas.setFillGray(0.90)
            canvas.setFont("Helvetica-Bold", 46)
            canvas.translate(width / 2, height / 2)
            canvas.rotate(40)
            canvas.drawCentredString(0, 0, plan.watermark)
        canvas.restoreState()

    story = [
        _pdf_paragraph(plan.title, title_style),
        Spacer(1, 4 * mm),
        _pdf_paragraph(plan.status, heading),
        PageBreak(),
    ]

    usable_width = A4[0] - 28 * mm
    for index, section in enumerate(plan.sections):
        story.append(_pdf_paragraph(section.title, heading))
        for paragraph in section.paragraphs:
            story.append(_pdf_paragraph(paragraph, body))

        for report_table in section.tables:
            if report_table.title:
                story.append(_pdf_paragraph(report_table.title, body))
            headers = [_pdf_paragraph(header, table_text) for header in report_table.headers]
            rows = [
                [_pdf_paragraph(cell, table_text) for cell in row]
                for row in _expanded_rows(report_table)
            ]
            data = [headers, *rows] if headers else rows
            if not data:
                story.append(_pdf_paragraph("No stored records.", body))
                continue
            columns = max(1, len(report_table.headers))
            table = Table(
                data,
                repeatRows=1,
                colWidths=[usable_width / columns] * columns,
                hAlign="LEFT",
            )
            table.setStyle(
                TableStyle(
                    [
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8E8E8")),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#AAAAAA")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 2),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
                        ("TOPPADDING", (0, 0), (-1, -1), 2),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                    ]
                )
            )
            story.extend([table, Spacer(1, 3 * mm)])

        if index < len(plan.sections) - 1:
            story.append(Spacer(1, 2 * mm))

    doc.build(story, onFirstPage=page, onLaterPages=page)
    return buffer.getvalue()


def render_docx(plan: ReportPlan) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    document.core_properties.title = plan.title
    document.core_properties.author = "SIH26035"

    if plan.watermark:
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = header.add_run(plan.watermark)
        run.bold = True
        run.font.size = Pt(18)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(plan.title)
    run.bold = True
    run.font.size = Pt(18)

    status = document.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = status.add_run(plan.status)
    run.bold = True
    run.font.size = Pt(13)

    document.add_page_break()

    for section_plan in plan.sections:
        document.add_heading(section_plan.title, level=1)
        for paragraph in section_plan.paragraphs:
            document.add_paragraph(paragraph)

        for report_table in section_plan.tables:
            if report_table.title:
                paragraph = document.add_paragraph()
                run = paragraph.add_run(report_table.title)
                run.bold = True

            expanded_rows = _expanded_rows(report_table)

            rows_count = len(expanded_rows) + 1
            columns_count = max(1, len(report_table.headers))
            table = document.add_table(
                rows=max(1, rows_count),
                cols=columns_count,
            )
            table.style = "Table Grid"

            for column, header in enumerate(report_table.headers):
                cell = table.cell(0, column)
                cell.text = header
                for run in cell.paragraphs[0].runs:
                    run.bold = True

            for row_index, row in enumerate(expanded_rows, start=1):
                for column, value in enumerate(row):
                    table.cell(row_index, column).text = value

            if not expanded_rows:
                table.cell(0, 0).text = "No stored records."

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_pair(context: dict) -> dict[str, tuple[bytes, str]]:
    plan = build_plan(context)
    return {
        "pdf": (render_pdf(plan), PDF_MIME),
        "docx": (render_docx(plan), DOCX_MIME),
    }
